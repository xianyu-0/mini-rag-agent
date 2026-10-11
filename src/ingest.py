import os, glob, chromadb
import docx
from pptx import Presentation
from pypdf import PdfReader
import openpyxl
from langchain_text_splitters import RecursiveCharacterTextSplitter
from .config import siliconflow_client, EMBED_MODEL


# ============================================================
# 文档解析器：把各种格式的文档转成纯文本
# ============================================================
def _load_docx(path: str) -> str:
    """提取 Word 文档的段落和表格"""
    doc = docx.Document(path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _load_pptx(path: str) -> str:
    """提取 PPT 每页的文本框"""
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text.strip():
                        slide_texts.append(text)
        if slide_texts:
            parts.append(f"[第 {i} 页]\n" + "\n".join(slide_texts))
    return "\n\n".join(parts)


def _load_pdf(path: str) -> str:
    """提取 PDF 每页文字"""
    reader = PdfReader(path)
    parts = []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"[第 {i} 页]\n{text}")
    return "\n\n".join(parts)


def _load_xlsx(path: str) -> str:
    """提取 Excel 每行数据"""
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"[工作表: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


# ============================================================
# 主加载函数：扫描 data/ 下所有支持的格式
# ============================================================
def load_docs(path="./data"):
    docs = []
    patterns = ["**/*.txt", "**/*.md", "**/*.docx", "**/*.pptx", "**/*.pdf", "**/*.xlsx"]
    files = []
    for pat in patterns:
        files.extend(glob.glob(f"{path}/{pat}", recursive=True))

    for f in files:
        try:
            ext = os.path.splitext(f)[1].lower()
            if ext in (".txt", ".md"):
                with open(f, encoding="utf-8") as fp:
                    text = fp.read()
            elif ext == ".docx":
                text = _load_docx(f)
            elif ext == ".pptx":
                text = _load_pptx(f)
            elif ext == ".pdf":
                text = _load_pdf(f)
            elif ext == ".xlsx":
                text = _load_xlsx(f)
            else:
                continue

            if text.strip():
                docs.append({"text": text, "source": os.path.basename(f)})
                print(f"  ✅ {os.path.basename(f)} ({len(text)} 字)")
            else:
                print(f"  ⚠️ {os.path.basename(f)} 提取为空（可能是扫描件/图片）")
        except Exception as e:
            print(f"  ❌ {os.path.basename(f)} 解析失败: {e}")

    print(f"\n📄 成功加载 {len(docs)} 篇文档")
    return docs


# ============================================================
# 文本切分
# ============================================================
def chunk_docs(docs, chunk_size=500, overlap=80):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size, chunk_overlap=overlap,
        separators=["\n## ", "\n### ", "\n\n", "\n", "。", " "]
    )
    chunks = []
    for d in docs:
        for i, c in enumerate(splitter.split_text(d["text"])):
            chunks.append({
                "id": f"{d['source']}::{i}",   # 幂等 ID，重复运行不会重复插入
                "text": c,
                "source": d["source"],
            })
    print(f"✂️ 切分成了 {len(chunks)} 个片段")
    return chunks


# ============================================================
# Embedding：调用硅基流动 bge-m3（分批，防止单次超限）
# ============================================================
def embed_texts(texts: list[str]) -> list[list[float]]:
    batch_size = 32
    all_embeddings = []
    for i in range(0, len(texts), batch_size):
        batch = texts[i:i + batch_size]
        resp = siliconflow_client.embeddings.create(
            model=EMBED_MODEL,
            input=batch,
        )
        all_embeddings.extend([d.embedding for d in resp.data])
        print(f"  ⏳ 已向量化 {min(i + batch_size, len(texts))}/{len(texts)}")
    return all_embeddings


# ============================================================
# 建库：写入 Chroma
# ============================================================
def build_index(chunks, persist_dir="./index"):
    client = chromadb.PersistentClient(path=persist_dir)
    print("🧮 开始调用硅基流动 bge-m3 向量化...")
    texts = [c["text"] for c in chunks]
    embeddings = embed_texts(texts)

    col = client.get_or_create_collection(
        name="docs",
        metadata={"hnsw:space": "cosine"},
    )
    col.upsert(
        ids=[c["id"] for c in chunks],
        documents=texts,
        embeddings=embeddings,
        metadatas=[{"source": c["source"]} for c in chunks],
    )
    print(f"✅ 入库 {len(chunks)} 个分块！")


# ============================================================
# 入口
# ============================================================
if __name__ == "__main__":
    docs = load_docs()
    if not docs:
        print("⚠️ data/ 文件夹里没有找到任何可解析的文档，请先放入文件。")
    else:
        chunks = chunk_docs(docs)
        build_index(chunks)